#! /usr/bin/python3

import openpyxl
import math

wb = openpyxl.Workbook()
sheet = wb["Sheet"]
my_pow = [i * i for i in range(1, 101)]
my_sqrt = [i ** (1 / 2) for i in range(1, 101)]
my_exp = [2**i for i in range(1, 101)]
my_log = [math.log2(i) for i in range(1, 101)]

sheet["A1"] = "Potęga"
sheet["B1"] = "Pierwiastek"
sheet["C1"] = "Exponanta"
sheet["D1"] = "Logarytm"

for i in enumerate(my_pow):
    sheet["A" + str(i[0] + 2)] = i[1]
for j in enumerate(my_sqrt):
    sheet["B" + str(j[0] + 2)] = j[1]
for k in enumerate(my_exp):
    sheet["C" + str(k[0] + 2)] = k[1]
for l in enumerate(my_log):
    sheet["D" + str(l[0] + 2)] = l[1]

refObjPow = openpyxl.chart.Reference(sheet, 1, 1, 1, 100)
seriesObjPow = openpyxl.chart.Series(refObjPow, title="Funkcja potęgowa")

refObjSqrt = openpyxl.chart.Reference(sheet, 2, 1, 2, 100)
seriesObjSqrt = openpyxl.chart.Series(
    refObjSqrt, title="Funkcja pierwiastek kwadratowy"
)

refObjExp = openpyxl.chart.Reference(sheet, 3, 1, 3, 100)
seriesObjExp = openpyxl.chart.Series(refObjExp, title="Funkcja wykładnicza")

refObjLog = openpyxl.chart.Reference(sheet, 4, 1, 4, 100)
seriesObjLog = openpyxl.chart.Series(refObjLog, title="Funkcja logarytmiczna")

chartObjPow = openpyxl.chart.BarChart()
chartObjPow.append(seriesObjPow)
chartObjSqrt = openpyxl.chart.BarChart()
chartObjSqrt.append(seriesObjSqrt)
chartObjExp = openpyxl.chart.BarChart()
chartObjExp.append(seriesObjExp)
chartObjLog = openpyxl.chart.BarChart()
chartObjLog.append(seriesObjLog)

# chartObj.drawing.top = 50
# chartObj.drawing.left = 100
# chartObj.drawing.width = 800
# chart.drawing.height = 500

sheet.add_chart(chartObjPow, "F1")
sheet.add_chart(chartObjSqrt, "F20")
sheet.add_chart(chartObjExp, "F40")
sheet.add_chart(chartObjLog, "F60")
wb.save("sampleChartCorrected.xlsx")
